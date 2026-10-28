"""요율표 엑셀 → 모바일 검색용 HTML 변환.

사용법: 이 폴더에 새 요율표 엑셀(.xlsx)을 넣고 `python build.py` 실행.
가장 최근에 수정된 엑셀 파일을 읽어 요율표.html(Artifact용)과 pwa/data.json(PWA용)을 만든다.

D열(생동/생산)은 엑셀 값을 쓰지 않고, H열 보험코드로 Druginform 서버 DB를 읽기 전용으로
조회해 자체생산/위탁생산(위탁업체명) 값으로 바꿔 넣는다. 서버에 접속할 수 없으면 직전에
받아둔 druginform_mfg.json 을 쓴다.

포장정보(의약품안전나라 제품정보의 '포장정보'와 같은 값)는 식약처 제품허가 상세 공공 API의
PACK_UNIT을 보험코드로 맞춰 넣는다. 전체 조회가 10분 남짓 걸려서 새 엑셀이 들어왔을 때만
다시 받고(druginform_pack.json 이 엑셀보다 오래됐을 때), 평소에는 저장된 사본을 쓴다.
API 키는 Druginform 서버의 것을 서버에서 읽어 쓰기만 한다(이 PC의 키는 미등록 상태).
"""
import sys
import json
import re
import subprocess
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
TEMPLATE = HERE / "template.html"
OUT = HERE / "요율표.html"
MFG_CACHE = HERE / "druginform_mfg.json"
PACK_CACHE = HERE / "druginform_pack.json"
PWA_DATA = HERE / "pwa" / "data.json"
SERVER = "root@1.234.80.8"

# 서버에서 실행할 읽기 전용 조회: 보험코드 → 위탁업체명(빈 문자열이면 자체생산).
# 아직 제조형태를 확인하지 못한 품목(mfg_checked_at NULL)은 넣지 않는다.
MFG_QUERY = r"""
import json, sqlite3
c = sqlite3.connect('file:/opt/drugapp/drugs.db?mode=ro', uri=True)
out = {}
for edi, cons, chk in c.execute("SELECT edi_code, consign_entp_name, mfg_checked_at FROM drugs "
                                "WHERE edi_code != '' AND (cancel_date IS NULL OR cancel_date = '')"):
    if not chk:
        continue
    for code in edi.split(','):
        if code.strip():
            out[code.strip()] = cons or ''
print(json.dumps(out, ensure_ascii=False))
"""

# 서버에서 실행할 식약처 제품허가 상세 API 전체 조회: 보험코드 → 포장단위.
# Druginform의 kdrug 클라이언트와 .env(API 키)를 읽어 쓰기만 하고, 서버에 아무것도 저장하지 않는다.
# 목록 API(getDrugPrdtPrmsnInq)에는 PACK_UNIT이 없어서 상세 API를 페이지 단위로 받는다.
PACK_QUERY = r"""
import json, sys, time
from kdrug import KdrugClient
c = KdrugClient.from_env()
out, page = {}, 1
while True:
    for attempt in range(4):
        try:
            rows = c._fetch(c.product_endpoint, {"numOfRows": 100, "pageNo": page, "type": "json"}, "product")
            break
        except Exception as e:
            if attempt == 3:
                raise
            time.sleep(5)
    if not rows:
        break
    for r in rows:
        pack = (r.get("PACK_UNIT") or "").strip()
        if not pack:
            continue
        cancelled = bool(r.get("CANCEL_DATE"))
        for code in (r.get("EDI_CODE") or "").split(","):
            code = code.strip()
            if code and (code not in out or not cancelled):
                out[code] = pack
    if page % 50 == 0:
        print(f"  {page}페이지, 보험코드 {len(out):,}건", file=sys.stderr, flush=True)
    page += 1
    time.sleep(0.2)
print(json.dumps(out, ensure_ascii=False))
"""


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, str):
        return re.sub(r"[ \t]+", " ", v).strip()
    return v


def clean_edi(v):
    """보험코드(H열). 엑셀에 숫자로 저장되면 앞자리 0이 빠지므로(051700030 → 51700030) 9자리로 되돌린다."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int):
        return f"{v:09d}"
    return clean(v)


def is_blank(row):
    return all(c in (None, "") for c in row)


def read_main(ws):
    rows, header_found = [], False
    for r in ws.iter_rows(values_only=True):
        if not header_found:
            header_found = bool(r) and clean(r[0]) == "성분명"
            continue
        if is_blank(r):
            continue
        row = [clean(c) for c in r[:8]]  # I열(특이사항)은 제외
        row += [""] * (8 - len(row))  # 끝쪽 빈 칸이 잘린 행도 8칸으로
        row[7] = clean_edi(r[7]) if len(r) > 7 else ""
        rows.append(row)
    # 머리글을 못 찾거나 행이 없으면 빈 요율표가 그대로 올라가므로 여기서 멈춘다
    if not header_found:
        raise SystemExit("첫 시트에서 A열이 '성분명'인 머리글 행을 찾지 못했습니다. 엑셀 첫 시트를 확인하세요.")
    if not rows:
        raise SystemExit("첫 시트 머리글 아래에 데이터가 없습니다. 엑셀 첫 시트를 확인하세요.")
    return rows


def read_special(ws):
    """'9월 신제품 및 재출시' / '9월 요율변동' / '9월 약가변동' 같은 구역으로 나눈다."""
    sections, cur = [], None
    for r in ws.iter_rows(values_only=True):
        if is_blank(r):
            continue
        filled = [c for c in r if c not in (None, "")]
        # 구역 제목: A열 한 칸에만 글자가 있는 행 (예: "9월 요율변동")
        if len(filled) == 1 and isinstance(r[0], str) and r[0].strip():
            cur = {"title": clean(r[0]), "header": None, "cols": [], "rows": []}
            sections.append(cur)
        elif cur is not None and cur["header"] is None:
            # 머리글 중간에 빈 칸이 있어도 열이 밀리지 않게, 머리글이 있는 칸의 위치를 기억한다
            cur["cols"] = [i for i, c in enumerate(r) if c not in (None, "")]
            cur["header"] = [clean(r[i]) for i in cur["cols"]]
        elif cur is not None:
            cur["rows"].append([
                (clean_edi if name == "보험코드" else clean)(r[i] if i < len(r) else None)
                for i, name in zip(cur["cols"], cur["header"])
            ])
    # 특이사항 열은 웹페이지에 싣지 않는다
    for s in sections:
        s.pop("cols", None)
        if s["header"] and "특이사항" in s["header"]:
            i = s["header"].index("특이사항")
            s["header"].pop(i)
            for row in s["rows"]:
                del row[i:i + 1]
    return sections


def load_mfg():
    try:
        res = subprocess.run(["ssh", SERVER, "python3", "-"], input=MFG_QUERY, capture_output=True,
                             text=True, encoding="utf-8", timeout=120, check=True)
        mfg = json.loads(res.stdout)
        if len(mfg) < 1000:  # 평소 2만 건 넘음 - 비정상적으로 적으면 사본을 덮어쓰지 않는다
            raise ValueError(f"제조형태가 {len(mfg)}건뿐입니다")
        MFG_CACHE.write_text(json.dumps(mfg, ensure_ascii=False), encoding="utf-8")
        return mfg, "서버 DB"
    except (OSError, subprocess.SubprocessError, ValueError) as e:
        if not MFG_CACHE.exists():
            raise SystemExit(f"Druginform 서버 DB 조회 실패, 저장된 사본도 없음: {e}")
        return json.loads(MFG_CACHE.read_text(encoding="utf-8")), f"저장된 사본(서버 조회 실패: {e})"


def load_pack(src):
    """새 엑셀이 들어왔을 때(사본이 엑셀보다 오래됐을 때)만 API로 다시 받는다."""
    if PACK_CACHE.exists() and PACK_CACHE.stat().st_mtime >= src.stat().st_mtime:
        return json.loads(PACK_CACHE.read_text(encoding="utf-8")), "저장된 사본"
    print("포장정보를 식약처 API에서 받는 중 (10분 남짓)...", flush=True)
    try:
        res = subprocess.run(["ssh", SERVER, "cd /opt/drugapp && venv/bin/python -"], input=PACK_QUERY,
                             stdout=subprocess.PIPE, stderr=sys.stderr, text=True, encoding="utf-8",
                             timeout=3600, check=True)
        pack = json.loads(res.stdout)
        if len(pack) < 1000:  # 평소 2만 건 넘음 - 비어 있는 결과를 사본으로 남기면 다음 엑셀까지 다시 안 받는다
            raise ValueError(f"포장정보가 {len(pack)}건뿐입니다")
        PACK_CACHE.write_text(json.dumps(pack, ensure_ascii=False), encoding="utf-8")
        return pack, "식약처 API"
    except (OSError, subprocess.SubprocessError, ValueError) as e:
        if PACK_CACHE.exists():
            return json.loads(PACK_CACHE.read_text(encoding="utf-8")), f"이전 사본(API 조회 실패: {e})"
        print(f"포장정보 조회 실패, 포장정보 없이 만듭니다: {e}")
        return {}, "없음"


def edi_key(edi):
    # 'D628800250'처럼 앞에 D가 붙은 보험코드는 같은 품목의 다른 유통채널 행 - D를 떼고 찾는다
    code = str(edi).strip()
    if code[:1] in ("D", "d") and code[1:].isdigit():
        code = code[1:]
    return code


def mfg_label(mfg, edi):
    code = edi_key(edi)
    if code not in mfg:
        return ""
    return f"위탁생산 · {mfg[code]}" if mfg[code] else "자체생산"


def apply_lookups(mfg, pack, data):
    """D열 → 제조형태로 교체, 맨 끝에 포장정보 열 추가."""
    hit_mfg = hit_pack = 0
    for r in data["rows"]:
        r[3] = mfg_label(mfg, r[7])
        r.append(pack.get(edi_key(r[7]), ""))
        hit_mfg += bool(r[3])
        hit_pack += bool(r[8])
    for s in data["special"]:
        h = s["header"] or []
        if "보험코드" not in h:
            continue
        j = h.index("보험코드")
        if "생동/생산" in h:
            i = h.index("생동/생산")
            for r in s["rows"]:
                r[i] = mfg_label(mfg, r[j])
        h.append("포장")
        for r in s["rows"]:
            r.append(pack.get(edi_key(r[j]), ""))
    return hit_mfg, hit_pack


def main():
    files = [p for p in HERE.glob("*.xlsx") if not p.name.startswith("~$")]
    if not files:
        raise SystemExit(f"{HERE} 폴더에 요율표 엑셀(.xlsx)이 없습니다.")
    src = max(files, key=lambda p: p.stat().st_mtime)
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    data = {
        "source": src.stem,
        "month": (re.search(r"\d{4}\.\d{2}", src.stem) or [src.stem])[0],
        "rows": read_main(wb.worksheets[0]),
        "special": read_special(wb.worksheets[1]) if len(wb.worksheets) > 1 else [],
    }
    mfg, mfg_src = load_mfg()
    pack, pack_src = load_pack(src)
    hit_mfg, hit_pack = apply_lookups(mfg, pack, data)
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = TEMPLATE.read_text(encoding="utf-8").replace("/*__DATA__*/null", payload)
    OUT.write_text(html, encoding="utf-8")
    # PWA(pwa/)용 데이터 - 같은 내용을 JSON으로. `python pwa/admin.py upload-data`로 Cloudflare에 올린다.
    PWA_DATA.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{src.name} → {OUT.name}: {len(data['rows'])}행, 특이사항 {len(data['special'])}구역, "
          f"{OUT.stat().st_size / 1e6:.1f}MB")
    print(f"제조형태({mfg_src}): {hit_mfg:,}/{len(data['rows']):,}행 매칭")
    print(f"포장정보({pack_src}): {hit_pack:,}/{len(data['rows']):,}행 매칭")


if __name__ == "__main__":
    main()
